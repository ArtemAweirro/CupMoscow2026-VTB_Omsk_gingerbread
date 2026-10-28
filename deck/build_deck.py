#!/usr/bin/env python3
"""Сборка презентации CupMoscow2026-ВТБ-ОмскийПряник.pdf: 8 слайдов + приложения.

Цифры берутся из модели и симуляции Разработчика 1, а не вписываются руками:
  out/model_values.json      итоги Excel (пишет model/check_model.py)
  model/…xlsx, «Допущения»   входы модели (оценки вариантов, сроки, допущения)
  out/sim_*.csv              симуляция (синтетика)
  out/0*.png                 графики для приложений
Поменяли допущение → ./run_all.sh → эта сборка → PDF с новыми цифрами.

Запуск:  .venv/bin/python deck/build_deck.py
Выход:   deck/build/deck.html (для правки и просмотра), deck/CupMoscow2026-ВТБ-ОмскийПряник.pdf
"""
import csv
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "deck"
BUILD = DECK / "build"
OUT = ROOT / "out"
XLSX = ROOT / "model" / "CupMoscow2026-ВТБ-ОмскийПряник.xlsx"
PDF = DECK / "CupMoscow2026-ВТБ-ОмскийПряник.pdf"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

# Тексты слайдов 2, 3, 6, 7 — черновик Разработчика 2 до текстов аналитика (31.10).
# Когда аналитик заменит тексты, поставить False: пометка «черновик» исчезнет.
DRAFT = True

VARIANTS = ["В1", "В2", "В3", "В4"]
VNAME = {
    "В1": "Правила + усиленная аутентификация",
    "В2": "ML-скоринг + сегментация",
    "В3": "Адаптивное трение",
    "В4": "Внешние участники",
}
VCOL = {"В1": "var(--v1)", "В2": "var(--v2)", "В3": "var(--v3)", "В4": "var(--v4)"}
VHEX = {"В1": "#eb6834", "В2": "#1baf7a", "В3": "#2a78d6", "В4": "#eda100"}

SRC_CBR = "Банк России, обзор ОБДС за II кв. 2026"


# ------------------------------------------------------------------ формат
def num(x, d=0):
    s = f"{x:,.{d}f}".replace(",", " ").replace(".", ",")
    return s.replace("-", "−")


def pct(x, d=0):
    return num(x * 100, d) + "%"


def signed(x, d=0):
    return ("+" if x > 0 else "") + num(x, d)


# ------------------------------------------------------------------ данные
def load():
    mv = json.loads((OUT / "model_values.json").read_text())
    wb = load_workbook(XLSX)
    ws = wb["Допущения"]
    inp, factors = {}, []
    for r in ws.iter_rows(min_row=6):
        key, name = r[0].value, r[1].value
        if isinstance(key, str) and key[:1] in "ACMORTV" and key[1:2].isdigit():
            inp[key] = {"name": name, "p": r[2].value, "b": r[3].value, "o": r[4].value,
                        "unit": r[5].value, "range": r[6].value, "status": r[7].value}
        if key is None and isinstance(name, str) and isinstance(r[2].value, float) and r[3].value in (1, 2, 3, 4, 5):
            factors.append((name, r[2].value, [r[3].value, r[4].value, r[5].value, r[6].value]))

    def rows(name):
        with open(OUT / name, encoding="utf-8") as f:
            return list(csv.DictReader(f))

    sim = {r["variant_short"]: {k: float(v) for k, v in r.items() if k not in ("variant_short", "variant")}
           for r in rows("sim_variants_metrics.csv")}
    curves = {}
    for r in rows("sim_curves.csv"):
        curves.setdefault(r["variant_short"], []).append(
            (float(r["legit_restricted_rate"]), float(r["loss_share"])))
    measures = rows("sim_measures_v3.csv")
    return mv, inp, factors, sim, curves, measures


# ------------------------------------------------------------------ каркас слайда
def slide(n, kicker, title, body, foot="", draft=False, lead=""):
    tag = '<span class="draft">черновик текста: заменит аналитик</span>' if (draft and DRAFT) else ""
    lead_html = f'<div class="lead">{lead}</div>' if lead else ""
    return f"""
<section class="slide">
  <div class="kicker">{kicker}{tag}</div>
  <h1>{title}</h1>{lead_html}
  <div class="body">{body}</div>
  <div class="foot"><div>{foot}</div><div class="num">{n}</div></div>
</section>"""


# ------------------------------------------------------------------ мини-графики (SVG/HTML)
def hbars(items, vmax, fmt, width_note=""):
    """items: [(label_html, value, color, bold)] → горизонтальные полосы."""
    out = ['<div class="bars">']
    for lab, v, col, bold in items:
        w = 0 if vmax == 0 else v / vmax * 78
        out.append(
            f'<div class="bar-row"><div class="lab" style="width:{width_note or "120px"}">{lab}</div>'
            f'<div class="track"><div class="fill" style="width:{w:.1f}%;background:{col}"></div>'
            f'<div class="val" style="left:calc({w:.1f}% + 6px);{"" if bold else "font-weight:normal"}">{fmt(v)}</div></div></div>')
    out.append("</div>")
    return "".join(out)


def curve_chart(curves, sim, w=470, h=330):
    """Потери против трения: ось X — доля законных операций с ограничением, ось Y — непредотвращенные потери."""
    ml, mr, mt, mb = 46, 24, 10, 40
    pw, ph = w - ml - mr, h - mt - mb
    xmax, ymax = 0.016, 0.8
    X = lambda x: ml + min(x, xmax) / xmax * pw
    Y = lambda y: mt + (1 - min(y, ymax) / ymax) * ph
    g = [f'<svg width="{w}" height="{h}" viewBox="0 0 {w} {h}" font-family="Arial" font-size="13">']
    for i in range(0, 5):
        yv = i * 0.2
        g.append(f'<line x1="{ml}" x2="{ml + pw}" y1="{Y(yv):.1f}" y2="{Y(yv):.1f}" stroke="#e6e5e0"/>')
        g.append(f'<text x="{ml - 6}" y="{Y(yv) + 4:.1f}" text-anchor="end" fill="#52514e">{int(yv * 100)}%</text>')
    for xv in (0, 0.004, 0.008, 0.012, 0.016):
        g.append(f'<text x="{X(xv):.1f}" y="{mt + ph + 18}" text-anchor="middle" fill="#52514e">{num(xv * 100, 1)}%</text>')
    g.append(f'<line x1="{ml}" x2="{ml + pw}" y1="{mt + ph}" y2="{mt + ph}" stroke="#8a8984"/>')
    g.append(f'<text x="{ml + pw / 2}" y="{h - 3}" text-anchor="middle" fill="#52514e">законные операции с ограничением</text>')
    for v in ["В1", "В2", "В4", "В3"]:
        pts = sorted((x, y) for x, y in curves[v] if x <= xmax)
        d = " ".join(f"{X(x):.1f},{Y(y):.1f}" for x, y in pts)
        g.append(f'<polyline points="{d}" fill="none" stroke="{VHEX[v]}" stroke-width="{3.2 if v == "В3" else 2.2}"/>')
        s = sim[v]
        g.append(f'<circle cx="{X(s["legit_restricted_rate"]):.1f}" cy="{Y(s["loss_share"]):.1f}" r="4.5" fill="{VHEX[v]}" stroke="#fff" stroke-width="1.5"/>')
    # подписи линий у правого края
    for v, dy in (("В1", -8), ("В2", -8), ("В4", -8), ("В3", -8)):
        pts = sorted((x, y) for x, y in curves[v] if x <= xmax)
        x, y = pts[-1]
        g.append(f'<text x="{X(x) - 2:.1f}" y="{Y(y) + dy:.1f}" text-anchor="end" fill="{VHEX[v]}" font-weight="bold">{v}</text>')
    g.append("</svg>")
    return "".join(g)


def waterfall(mv, w=530, h=236):
    b = mv["scenarios"]["b"]
    steps = [("Возмещения", b["avoided"], "+"), ("Ограничения", b["fp_saving"], "+"),
             ("Удержание", b["retention"], "+"), ("Эффект", b["gross"], "="), ("Расходы", -b["opex"], "+"),
             ("Чистый", b["net"], "=")]
    ml, mt, mb = 34, 22, 26
    ph = h - mt - mb
    vmax = b["gross"] * 1.05
    bw, gap = (w - ml) / len(steps) * 0.72, (w - ml) / len(steps) * 0.28
    Y = lambda v: mt + (1 - v / vmax) * ph
    g = [f'<svg width="{w}" height="{h}" viewBox="0 0 {w} {h}" font-family="Arial" font-size="13">']
    run = 0
    for i, (lab, v, kind) in enumerate(steps):
        x = ml + i * (bw + gap)
        if kind == "=":
            y0, y1, col = 0, v, "#52514e"
            run = v
        else:
            y0, y1 = run, run + v
            col = "#2a78d6" if v >= 0 else "#eb6834"
            run += v
        top, bot = Y(max(y0, y1)), Y(min(y0, y1))
        g.append(f'<rect x="{x:.1f}" y="{top:.1f}" width="{bw:.1f}" height="{max(bot - top, 1):.1f}" fill="{col}" rx="2"/>')
        txt = num(v) if kind == "=" else signed(v)
        g.append(f'<text x="{x + bw / 2:.1f}" y="{top - 5:.1f}" text-anchor="middle" font-weight="bold" fill="#0b0b0b">{txt}</text>')
        words = lab.split(" ")
        for j, wd in enumerate(words[:2]):
            g.append(f'<text x="{x + bw / 2:.1f}" y="{h - mb + 17 + j * 15}" text-anchor="middle" fill="#52514e">{wd}</text>')
    g.append(f'<line x1="{ml}" x2="{w}" y1="{Y(0):.1f}" y2="{Y(0):.1f}" stroke="#8a8984"/>')
    g.append("</svg>")
    return "".join(g)


# ------------------------------------------------------------------ макеты экранов
def phone(screen):
    head = '<div class="island"></div><div class="status"><span>9:41</span><span>●●● ▮</span></div>'
    if screen == 1:
        body = """
<div class="bar"><span class="back">‹</span>Перевод по СБП</div>
<div class="scr">
  <div><div class="amt">150 000 ₽</div><div class="rcp">Новый получатель · Игорь С. · другой банк</div></div>
  <div class="warn"><div class="wt">Вам сейчас звонили от имени банка, полиции или госоргана?</div>
  Банк никогда не просит переводить деньги на «безопасный счёт». Если вас торопят по телефону, положите трубку.</div>
</div>
<div class="bottom"><div class="btn ds">Меня просят перевести — отменить</div><div class="btn p">Нет, это мой перевод</div></div>"""
    elif screen == 2:
        body = """
<div class="bar"><span class="back">‹</span>Подтвердите перевод</div>
<div class="scr">
  <div><div class="amt">48 000 ₽</div><div class="rcp">Новый получатель · Олег В.</div></div>
  <div class="info">Вы вчера сменили номер телефона. Подтвердите, что перевод делаете вы.</div>
  <svg width="64" height="64" viewBox="0 0 64 64" style="align-self:center;margin-top:6px">
    <rect x="3" y="3" width="58" height="58" rx="16" fill="none" stroke="#2a78d6" stroke-width="3"/>
    <circle cx="23" cy="26" r="3" fill="#2a78d6"/><circle cx="41" cy="26" r="3" fill="#2a78d6"/>
    <path d="M32 26 v9 h-3" fill="none" stroke="#2a78d6" stroke-width="3" stroke-linecap="round"/>
    <path d="M22 42 q10 8 20 0" fill="none" stroke="#2a78d6" stroke-width="3" stroke-linecap="round"/></svg>
  <div class="center tiny">Код из SMS не нужен. Никому не сообщайте коды.</div>
</div>
<div class="bottom"><div class="btn p">Подтвердить по биометрии</div><div class="btn s">Это не я — заблокировать</div></div>"""
    elif screen == 3:
        body = """
<div class="bar"><span class="back">‹</span>Перевод на паузе</div>
<div class="scr">
  <div class="timer">3:42:15</div>
  <div class="info">Перевод 300 000 ₽ похож на схему «безопасный счёт». Мы придержали его, чтобы вы могли спокойно проверить.
  Вам позвонит сотрудник банка с номера, указанного на обороте карты.</div>
  <div class="tiny">Деньги остаются на вашем счёте. Банк не просит назвать коды.</div>
</div>
<div class="bottom"><div class="btn d">Меня обманывают — отменить</div><div class="btn s">Это я, продолжить после звонка</div>
<div class="tiny center">«Продолжить» недоступно, если сработал признак по 161-ФЗ</div></div>"""
    else:
        body = """
<div class="bar"><span class="back">‹</span>Защита близких</div>
<div class="scr">
  <div class="b" style="font-size:16px">Пусть крупные переводы подтверждает человек, которому вы доверяете</div>
  <div class="field"><div class="fl">Защитник</div>Анна, дочь</div>
  <div class="field"><div class="fl">Подтверждение переводов от</div>50 000 ₽</div>
  <div class="info">Бесплатно. Защитник отвечает в приложении; нет ответа за 12 часов — перевод отменяется. Отключение — через 24 часа.</div>
</div>
<div class="bottom"><div class="btn p">Подключить</div><div class="btn s">Не сейчас</div></div>"""
    return f'<div class="phone">{head}{body}<div class="home"></div></div>'


def phone_small(screen, scale):
    w, h = 270 * scale, 556 * scale
    return (f'<div style="width:{w:.0f}px;height:{h:.0f}px;flex:0 0 auto">'
            f'<div class="scaled" style="transform:scale({scale})">{phone(screen)}</div></div>')


# ------------------------------------------------------------------ схемы
def links_svg(w, h, paths, markers=True):
    out = [f'<svg class="links" width="{w}" height="{h}" viewBox="0 0 {w} {h}">',
           '<defs><marker id="ar" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="7" markerHeight="7" orient="auto-start-reverse">'
           '<path d="M0,0 L10,5 L0,10 z" fill="#52514e"/></marker>'
           '<marker id="arn" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="7" markerHeight="7" orient="auto-start-reverse">'
           '<path d="M0,0 L10,5 L0,10 z" fill="#104281"/></marker></defs>']
    for pts, style in paths:
        d = "M" + " L".join(f"{x:.1f},{y:.1f}" for x, y in pts)
        if style == "reg":
            out.append(f'<path d="{d}" fill="none" stroke="#104281" stroke-width="1.6" stroke-dasharray="5 4" marker-end="url(#arn)"/>')
        elif style == "loop":
            out.append(f'<path d="{d}" fill="none" stroke="#8a8984" stroke-width="1.5" stroke-dasharray="4 4" marker-end="url(#ar)"/>')
        else:
            out.append(f'<path d="{d}" fill="none" stroke="#52514e" stroke-width="1.5" marker-end="url(#ar)"/>')
    out.append("</svg>")
    return "".join(out)


def elbow(a, b, gap=0):
    """Ортогональная связь от правого края a к левому краю b (a, b = x, y, w, h)."""
    ax, ay = a[0] + a[2], a[1] + a[3] / 2
    bx, by = b[0] - gap, b[1] + b[3] / 2
    if abs(ay - by) < 1:
        return [(ax, ay), (bx, by)]
    mx = (ax + bx) / 2
    return [(ax, ay), (mx, ay), (mx, by), (bx, by)]


def down(a, b):
    ax, ay = a[0] + a[2] / 2, a[1] + a[3]
    return [(ax, ay), (ax, b[1])]


def swimlane():
    lanes = ["Клиент", "Антифрод-система", "Фрод-мониторинг и контакт-центр", "Безопасность и комплаенс", "ИТ и данные"]
    stages = ["1. Операция", "2. Сбор признаков", "3. Оценка риска", "4. Решение", "5. Действие клиента", "6. Разбор", "7. Обучение"]
    LW, top, lh = 136, 24, 52
    cw = (1200 - LW) / len(stages)
    bw, bh = cw - 14, lh - 10

    def at(lane, col):
        return (LW + col * cw + 7, top + lane * lh + 5, bw, bh)

    boxes = {
        "op": (0, 0, "Перевод СБП или со счёта", "client"),
        "data": (4, 1, "Данные; часть — с задержкой до 24 ч", "plain"),
        "feat": (1, 1, "Сбор признаков в моменте", ""),
        "score": (1, 2, "Скоринг риска и сценария", ""),
        "reg": (3, 2, "Регуляторный слой 161-ФЗ", "reg"),
        "engine": (1, 3, "Движок: мера по риску", "main"),
        "act": (0, 4, "Предупреждение, биометрия, пауза", "client"),
        "call": (2, 4, "Звонок по скрипту", ""),
        "res": (0, 5, "Платёж прошёл или отменён", "client"),
        "review": (2, 5, "Разбор спорных кейсов", ""),
        "obds": (3, 5, "Заявления об ОБДС", "plain"),
        "learn": (4, 6, "Переобучение, A/B, KPI", "plain"),
    }
    pos = {k: at(l, c) for k, (l, c, _, _) in boxes.items()}
    h_total = top + len(lanes) * lh + 18
    html = [f'<div class="canvas" style="width:1200px;height:{h_total}px">']
    for i, s in enumerate(stages):
        html.append(f'<div class="stage-label" style="left:{LW + i * cw:.0f}px;top:0;width:{cw:.0f}px">{s}</div>')
    for i, l in enumerate(lanes):
        y = top + i * lh
        bg = "background:#f8f9fb;" if i % 2 == 0 else ""
        html.append(f'<div class="lane-bg" style="left:0;top:{y}px;width:1200px;height:{lh}px;{bg}"></div>')
        html.append(f'<div class="lane-label" style="left:0;top:{y}px;width:{LW - 6}px;height:{lh}px">{l}</div>')
    paths = [
        (elbow(pos["op"], pos["feat"]), ""),
        ([(pos["data"][0] + bw / 2, pos["data"][1]), (pos["data"][0] + bw / 2, pos["feat"][1] + bh)], ""),
        (elbow(pos["feat"], pos["score"]), ""),
        (elbow(pos["score"], pos["engine"]), ""),
        (elbow(pos["reg"], pos["engine"]), "reg"),
        (elbow(pos["engine"], pos["act"]), ""),
        (elbow(pos["engine"], pos["call"]), ""),
        ([(pos["call"][0] + bw / 2, pos["call"][1]), (pos["call"][0] + bw / 2, pos["act"][1] + bh)], ""),
        (elbow(pos["act"], pos["res"]), ""),
        (elbow(pos["call"], pos["review"]), ""),
        ([(pos["review"][0] + bw / 2, pos["review"][1] + bh), (pos["review"][0] + bw / 2, pos["obds"][1])], ""),
        (elbow(pos["obds"], pos["learn"]), ""),
    ]
    # обратная связь: обучение → скоринг (по нижнему краю и зазору между колонками)
    lx, ly, _, _ = pos["learn"]
    sx, sy, _, _ = pos["score"]
    ybot = top + len(lanes) * lh + 10
    gx = LW + 2 * cw + 2
    paths.append(([(lx + bw / 2, ly + bh), (lx + bw / 2, ybot), (gx, ybot), (gx, sy + bh + 4), (sx + 22, sy + bh + 4),
                   (sx + 22, sy + bh)], "loop"))
    html.append(links_svg(1200, h_total, paths))
    for k, (l, c, text, cls) in boxes.items():
        x, y, w, h = pos[k]
        html.append(f'<div class="box {cls}" style="left:{x:.0f}px;top:{y:.0f}px;width:{w:.0f}px;height:{h:.0f}px">{text}</div>')
    html.append(f'<div class="small muted" style="position:absolute;left:{gx + 8:.0f}px;top:{ybot - 9}px;background:#fff;padding:0 4px">обратная связь: новая версия модели</div>')
    html.append("</div>")
    return "".join(html), h_total


ARCH = [
    # key, x, y, w, h, title, purpose, input, output, cls
    ("data", 0, 100, 240, 190, "Сбор и обогащение данных", "признаки в моменте + пакетная догрузка",
     "транзакция, устройство, сессия, номер, получатель", "признаки и флаг «данные неполные»", ""),
    ("reg", 290, 0, 280, 118, "Регуляторный слой 161-ФЗ", "обязательные меры, не обучается",
     "12 признаков ОД-2506, база Банка России", "приостановка 2 дня или отказ", "reg"),
    ("score", 290, 128, 280, 118, "Скоринг риска", "риск и тип сценария (А / Б)",
     "признаки", "оценка риска 0–1, тип сценария", ""),
    ("seg", 290, 256, 280, 118, "Сегментация и профиль", "индивидуальные пороги",
     "история, возраст, «Защита близких»", "порог клиента", ""),
    ("engine", 620, 90, 260, 210, "Движок решений", "мера по уровню риска, а не «пропустить или блок»",
     "оценка, тип, порог, регуляторный сигнал",
     "пропустить → предупредить → биометрия → пауза → звонок → отказ", "main"),
    ("measures", 930, 30, 270, 118, "Клиентские меры", "понятный шаг для клиента",
     "мера движка", "подтверждение, отмена, «меня обманывают»", ""),
    ("ops", 930, 190, 270, 118, "Операционный контур", "звонки, чат-бот, разбор",
     "кейсы высокого риска", "решение по кейсу, заявление об ОБДС", ""),
    ("learn", 290, 398, 570, 76, "Обратная связь и обучение", "переобучение модели, A/B-тесты",
     "подтверждения, заявления об ОБДС, итоги разбора", "новая версия модели", ""),
    ("kpi", 900, 398, 300, 76, "Мониторинг KPI", "пороги",
     "потери, NPS, ограничения", "новые пороги", ""),
]


def architecture():
    pos = {k: (x, y, w, h) for k, x, y, w, h, *_ in ARCH}
    H = 476
    html = [f'<div class="canvas" style="width:1200px;height:{H}px">']
    p = pos
    cx = lambda k: p[k][0] + p[k][2] / 2
    paths = [
        (elbow(p["data"], p["reg"]), "reg"), (elbow(p["data"], p["score"]), ""), (elbow(p["data"], p["seg"]), ""),
        (elbow(p["reg"], p["engine"]), "reg"), (elbow(p["score"], p["engine"]), ""), (elbow(p["seg"], p["engine"]), ""),
        (elbow(p["engine"], p["measures"]), ""), (elbow(p["engine"], p["ops"]), ""),
        (down(p["ops"], p["kpi"]), ""),
        ([(p["kpi"][0], p["kpi"][1] + p["kpi"][3] / 2), (p["learn"][0] + p["learn"][2], p["learn"][1] + p["learn"][3] / 2)], ""),
        ([(cx("seg"), p["learn"][1]), (cx("seg"), p["seg"][1] + p["seg"][3])], "loop"),
        ([(cx("engine"), p["learn"][1]), (cx("engine"), p["engine"][1] + p["engine"][3])], "loop"),
    ]
    html.append(links_svg(1200, H, paths))
    for k, x, y, w, h, t, pu, i, o, cls in ARCH:
        if h < 100:
            inner = (f'<div class="t">{t}</div><div><span class="k">Назначение:</span> {pu}. '
                     f'<span class="k">Вход:</span> {i}. <span class="k">Выход:</span> {o}.</div>')
        else:
            inner = (f'<div class="t">{t}</div><div><span class="k">Назначение:</span> {pu}</div>'
                     f'<div><span class="k">Вход:</span> {i}</div><div><span class="k">Выход:</span> {o}</div>')
        html.append(f'<div class="block {cls}" style="left:{x}px;top:{y}px;width:{w}px;height:{h}px">{inner}</div>')
    html.append(f'<div class="small muted" style="position:absolute;left:{cx("seg") + 8:.0f}px;top:{p["learn"][1] - 22}px">новая модель</div>')
    html.append(f'<div class="small muted" style="position:absolute;left:{cx("engine") + 8:.0f}px;top:{p["learn"][1] - 22}px">новые пороги</div>')
    html.append("</div>")
    return "".join(html)


# ------------------------------------------------------------------ слайды
def s1(mv):
    b = mv["scenarios"]["b"]
    return f"""
<section class="slide cover">
  <div class="inner">
    <div class="tag">Кейс ВТБ «Деньги под защитой» · Cup Moscow 2026 · секция ВТБ</div>
    <div class="name">Адаптивное трение:<br>защитить деньги, не мешая платить</div>
    <div class="sub">Мера защиты соответствует риску перевода: предупреждение и биометрия вместо блокировки там, где риск невысок; пауза и звонок — там, где клиента обманывают</div>
  </div>
  <div class="facts">
    <div class="fact"><div class="v">−{num(b["saved"])} млн ₽</div><div class="l">потерь клиентов в год (СБП и переводы со счёта)</div></div>
    <div class="fact"><div class="v">−{pct(1 - b["restr_after"] / b["restr_before"])}</div><div class="l">ошибочных ограничений законных переводов</div></div>
    <div class="fact"><div class="v">{num(mv["costs"]["total"], 1)} млн ₽ · 6 мес.</div><div class="l">внедрение в рамках кейса (≤ 50 млн ₽, ≤ 6 мес.)</div></div>
  </div>
  <div class="team"><span>Команда <b>«Омский пряник»</b></span><span>Базовый сценарий; расчёты — в Excel-приложении с тем же именем</span></div>
</section>"""


def s2(mv, inp):
    ch = [("СБП", inp["M7"]["b"], inp["M8"]["b"], True), ("Счета (ДБО)", inp["M5"]["b"], inp["M6"]["b"], True),
          ("Карты", inp["M9"]["b"], inp["M10"]["b"], False), ("Без открытия счёта", inp["M13"]["b"], inp["M14"]["b"], False),
          ("Эл. кошельки", inp["M11"]["b"], inp["M12"]["b"], False)]
    tot_s = sum(c[1] for c in ch)
    tot_n = sum(c[2] for c in ch)
    foc_s = ch[0][1] + ch[1][1]
    foc_n = ch[0][2] + ch[1][2]
    bars_sum = hbars([(c[0], c[1], "var(--blue)" if c[3] else "#c9c8c3", c[3]) for c in ch], max(c[1] for c in ch),
                     lambda v: num(v, 2), "128px")
    bars_n = hbars([(c[0], c[2], "var(--blue)" if c[3] else "#c9c8c3", c[3]) for c in ch], max(c[2] for c in ch),
                   lambda v: num(v, 0), "128px")
    title = f"СБП и переводы со счёта — {pct(foc_s / tot_s)} похищенного у физлиц: на них и фокусируемся"
    body = f"""
<div class="row" style="height:100%">
  <div style="flex:0 0 380px" class="col">
    <div><div class="h3">Сумма хищений у физлиц<span class="sub">млрд ₽, II кв. 2026</span></div>{bars_sum}</div>
    <div><div class="h3">Число операций<span class="sub">тыс., II кв. 2026</span></div>{bars_n}</div>
    <div class="note">Фокус — <b class="blue">{pct(foc_s / tot_s)} денег</b> при {pct(foc_n / tot_n)} операций: средний чек ≈{num(foc_s / foc_n * 1000)} тыс. ₽ против ≈{num(ch[2][1] / ch[2][2] * 1000, 1)} тыс. ₽ у карт.</div>
  </div>
  <div style="flex:0 0 400px">
    <div class="h3">Почему банк теряет деньги и мешает клиентам</div>
    <div class="why">
      <div class="lvl"><div class="n">1</div><div class="tx">Клиенты теряют деньги, а законные платежи останавливаются</div></div>
      <div class="arrow">почему?</div>
      <div class="lvl"><div class="n">2</div><div class="tx">{num(inp["M3"]["b"] / 1000, 2)} млрд ₽ ОБДС за квартал по рынку; ошибочные ограничения — обращения, отказ от операций, отток</div></div>
      <div class="arrow">почему?</div>
      <div class="lvl"><div class="n">3</div><div class="tx">Перевод под обманом выглядит как обычный: клиент сам входит и подтверждает</div></div>
      <div class="arrow">почему?</div>
      <div class="lvl"><div class="n">4</div><div class="tx">Верный код ≠ добровольность; данные приходят с задержкой до суток; пороги одни для всех</div></div>
      <div class="arrow">почему?</div>
      <div class="lvl"><div class="n">5</div><div class="tx">Одна мера на любой риск — «пропустить или заблокировать»; нет независимого участника</div></div>
    </div>
  </div>
  <div style="flex:1" class="card key">
    <div class="h3 navy">Фокус решения</div>
    <div class="small"><b>Операции:</b> переводы физлиц в мобильном и интернет-банке — СБП и со счёта</div>
    <div class="small" style="margin-top:7px"><b>Сценарии:</b> А — «безопасный счёт» (социальная инженерия); Б — захват доступа: смена номера или устройства, затем перевод новому получателю</div>
    <div class="small" style="margin-top:7px"><b>Сегменты риска:</b> 55+; после смены номера или устройства; новый получатель; нетипичная сумма</div>
    <div class="small" style="margin-top:7px"><b>Каналы выявления:</b> оценка в момент перевода; пакетный анализ до 24 ч; звонок банка при высоком риске</div>
    <div class="small muted" style="margin-top:9px"><b>Не берём:</b> карты и банкоматы (этап 2), юрлица, электронные кошельки</div>
  </div>
</div>"""
    return slide(2, "Проблема и фокус", title, body, f"Источник: {SRC_CBR}; материалы кейса.", draft=True)


RISK = {
    "В1": ("обходится социнженерией", "базовый слой правил"),
    "В2": ("при сомнении — блок", "скоринг войдёт в В3"),
    "В3": ("настройка порогов и UX", "пилот 20% трафика, A/B"),
    "В4": ("зависит от партнёров", "этап 2 после В3"),
}
FACTOR_SHORT = {
    "Качество выявления (полнота)": "Качество выявления",
    "Ошибочные ограничения (ложные срабатывания)": "Ошибочные ограничения",
    "Влияние на NPS и удобство": "NPS и удобство",
    "Стоимость внедрения": "Стоимость внедрения",
    "Срок (укладывается ли в 6 мес.)": "Срок ≤ 6 мес.",
    "Соответствие 161-ФЗ": "Соответствие 161-ФЗ",
    "Устойчивость к задержке данных до суток": "Задержка данных до суток",
    "Риски внедрения": "Риски внедрения",
}


def s3(mv, inp, factors, sim):
    sc = mv["variant_score"]
    cost = mv["variant_cost"]
    term = {"В1": inp["V1t"]["b"], "В2": inp["V2t"]["b"], "В3": inp["C2"]["b"], "В4": inp["V4t"]["b"]}
    hdr = "".join(f'<th class="c{" hl" if v == "В3" else ""}" style="width:132px"><span class="vdot" style="background:{VCOL[v]}"></span>{v}<br>'
                  f'<span style="font-weight:normal">{VNAME[v]}</span></th>' for v in VARIANTS)
    rows = []
    for i, (name, w, vals) in enumerate(factors):
        bold = ' class="b"' if i < 3 else ""
        cells = "".join(f'<td class="c{" hl" if v == "В3" else ""}">{x}</td>' for v, x in zip(VARIANTS, vals))
        rows.append(f'<tr><td{bold}>{FACTOR_SHORT.get(name, name)}</td><td class="num"{bold}>{pct(w)}</td>{cells}</tr>')
    rows.append('<tr class="total"><td>Итого, из 5</td><td></td>' +
                "".join(f'<td class="c{" hl" if v == "В3" else ""}">{num(sc[v], 2)}</td>' for v in VARIANTS) + "</tr>")
    rows.append('<tr><td>Стоимость, срок</td><td></td>' +
                "".join(f'<td class="c{" hl" if v == "В3" else ""}">{num(cost[v], 1 if v == "В3" else 0)} млн · {num(term[v])} мес.</td>' for v in VARIANTS) + "</tr>")
    rows.append('<tr><td>Главный риск</td><td></td>' +
                "".join(f'<td class="{"hl" if v == "В3" else ""}">{RISK[v][0]}</td>' for v in VARIANTS) + "</tr>")
    rows.append('<tr><td>Следующий шаг</td><td></td>' +
                "".join(f'<td class="{"hl" if v == "В3" else ""}">{RISK[v][1]}</td>' for v in VARIANTS) + "</tr>")
    table = (f'<table class="tight"><tr><th>Фактор</th><th class="num">Вес</th>{hdr}</tr>{"".join(rows)}</table>')
    lab = lambda v: f'<span class="vdot" style="background:{VCOL[v]}"></span>{v}'
    eq_fr = sim["В1"]["legit_restricted_rate"]
    bars1 = hbars([(lab(v), sim[v]["loss_share"], VCOL[v], v == "В3") for v in VARIANTS], 0.5, lambda x: pct(x), "48px")
    bars2 = hbars([(lab(v), sim[v]["restricted_rate_at_equal_loss"], VCOL[v], v == "В3") for v in VARIANTS], 0.014,
                  lambda x: pct(x, 2), "48px")
    best = sim["В3"]
    title = (f"Выбираем адаптивное трение (В3): {num(sc['В3'], 1)} из 5 и наименьшие потери "
             f"при том же трении")
    body = f"""
<div class="row" style="height:100%">
  <div style="flex:0 0 790px">{table}
    <div class="note" style="margin-top:6px">Выделены приоритетные факторы (60% веса): они отражают суть кейса — потери и удобство клиента. Оценки 1–5 — команда, уточнены симуляцией.</div>
  </div>
  <div style="flex:1" class="col">
    <div class="card soft">
      <div class="h3">Потери при равном трении</div>
      <div class="note">непредотвращённые потери, % суммы попыток; у всех {pct(eq_fr, 2)} законных операций с ограничением</div>
      {bars1}
    </div>
    <div class="card soft">
      <div class="h3">Трение при равных потерях</div>
      <div class="note">законные операции с ограничением, если удерживать потери на 35%</div>
      {bars2}
    </div>
    <div class="note">В3 гасит {pct(best["signals_cleared_share"])} сигналов предупреждением или биометрией — они не останавливают платёж. Синтетические данные: иллюстрация механизма, не прогноз.</div>
  </div>
</div>"""
    return slide(3, "Варианты решения", title, body,
                 "Симуляция на синтетических данных — Приложение 4; стоимость В3 — смета в Excel, лист «Затраты».", draft=True)


def s4():
    lane, hl = swimlane()
    scen = f"""
<div class="row" style="margin-top:6px;gap:20px">
  <div class="card" style="flex:1;display:flex;gap:14px;padding:10px 12px">
    {phone_small(2, 0.3)}
    <div>
      <div class="h3">«Стажёр»: обычная покупка не срывается</div>
      <div class="small">Вчера сменил номер, сегодня переводит 48 тыс. ₽ новому получателю → риск средний, похоже на захват доступа →
      <b>биометрия вместо блокировки</b> → платёж проходит за секунды, продавец не ждёт.</div>
      <div class="small muted" style="margin-top:6px">Мошенник без лица владельца биометрию не пройдёт.</div>
    </div>
  </div>
  <div class="card" style="flex:1;display:flex;gap:10px;padding:10px 12px">
    {phone_small(3, 0.3)}{phone_small(4, 0.3)}
    <div>
      <div class="h3">«Бабушка»: деньги остаются у клиента</div>
      <div class="small">Звонок «из спецслужб», перевод 300 тыс. ₽ на «безопасный счёт» → риск высокий, сценарий А →
      <b>пауза и звонок банка</b> → клиент отменяет перевод → банк предлагает «Защиту близких».</div>
      <div class="small muted" style="margin-top:6px">Для снятия наличных — та же мера: пауза и звонок.</div>
    </div>
  </div>
</div>"""
    body = lane + scen
    return slide(4, "Целевой процесс и клиентский путь",
                 "Банк отвечает на риск ступенчато: обычный клиент подтверждает за секунды, жертву обмана останавливают пауза и звонок",
                 body, "Пунктир: регуляторный путь 161-ФЗ (приостановка до 2 дней, период охлаждения) — модель его не отменяет. Экраны — Приложение 5.")


def s5():
    biz = """
<div class="row small" style="margin-top:10px;gap:10px;align-items:center">
  <div class="b" style="flex:0 0 auto">Бизнес-инструменты:</div>
  <div class="card soft" style="padding:5px 10px">скрипт звонка при паузе</div>
  <div class="card soft" style="padding:5px 10px">методика порогов по сегментам</div>
  <div class="card soft" style="padding:5px 10px">регламент разбора и возврата</div>
  <div class="card soft" style="padding:5px 10px">тексты предупреждений</div>
  <div class="card soft" style="padding:5px 10px">обучение контакт-центра</div>
</div>"""
    return slide(5, "Архитектура и блоки",
                 "Девять блоков встраиваются в антифрод банка; регуляторный слой работает отдельно и не переопределяется моделью",
                 architecture() + biz,
                 "Данные с задержкой до суток: пауза даёт время догрузить их и перепроверить риск. Назначение, вход и выход бизнес-блоков — Приложение 6.")


def s6(mv, sim, curves):
    s = mv["scenarios"]
    P, B, O = s["p"], s["b"], s["o"]
    rows = [
        ("а", "Похищенные средства клиентов", "млн ₽/год", "попытки × (1 − полнота)",
         num(B["loss_before"]), [num(x["loss_after"]) for x in (P, B, O)]),
        ("б", "Корректно выявлено мошеннических операций", "тыс./год", "попытки × полнота",
         num(B["detected_before"]), [num(x["detected_after"]) for x in (P, B, O)]),
        ("", "  полнота выявления по сумме", "%", "допущение", pct(0.8), [pct(x["recall_after"]) for x in (P, B, O)]),
        ("в", "Ошибочные ограничения", "% законных переводов", "сигналы × доля ставших ограничением",
         pct(B["restr_before"], 3), [pct(x["restr_after"], 3) for x in (P, B, O)]),
        ("г", "NPS банка", "п., изменение", "затронутые клиенты × разрыв NPS", "—",
         [signed(x["nps_delta"], 1) for x in (P, B, O)]),
        ("д", "Попытки мошенничества", "изменение", "эффект сдерживания, только оптимист", "—", ["0%", "0%", "−3%"]),
    ]
    tr = []
    for k, name, unit, f, was, vals in rows:
        nm = f'<b>{k})</b> {name}' if k else f'<span class="muted">{name.strip()}</span>'
        tr.append(f'<tr><td>{nm}<br><span class="muted">{unit}</span></td><td class="muted">{f}</td><td class="num">{was}</td>'
                  f'<td class="num">{vals[0]}</td><td class="num hl b">{vals[1]}</td><td class="num">{vals[2]}</td></tr>')
    table = ('<table><tr><th style="width:250px">Показатель</th><th>Как считаем</th><th class="num">Было</th>'
             '<th class="num">Пессимист</th><th class="num hl">База</th><th class="num">Оптимист</th></tr>' + "".join(tr) + "</table>")
    title = (f"База: потери клиентов −{num(B['saved'])} млн ₽ в год, ошибочных ограничений "
             f"−{pct(1 - B['restr_after'] / B['restr_before'])}, NPS +{num(B['nps_delta'], 1)} п.")
    body = f"""
<div class="row" style="height:100%">
  <div style="flex:0 0 700px">{table}
    <div class="note" style="margin-top:8px"><b>Сохранённые средства клиентов</b> ({num(B["saved"])} млн ₽) — KPI клиента, не доход банка.
    <b>Сигнал ≠ ограничение:</b> предупреждение и биометрия не останавливают платёж, поэтому ограничений меньше при большем числе сигналов.</div>
  </div>
  <div style="flex:1">
    <div class="h3">Почему ограничений меньше: потери против трения</div>
    <div class="note" style="margin-bottom:4px">непредотвращённые потери, % суммы попыток; точка — рабочий порог; ниже и левее лучше</div>
    {curve_chart(curves, sim)}
    <div class="note">Синтетические данные, задержка до суток учтена: иллюстрация механизма, не прогноз.</div>
  </div>
</div>"""
    return slide(6, "Влияние на KPI", title, body,
                 "Допущения и сценарии — Excel, листы «Допущения», «KPI», «Сценарии»; база рынка: " + SRC_CBR + ".", draft=True)


def s7(mv, inp):
    s = mv["scenarios"]
    B = s["b"]
    c = mv["costs"]
    cost_rows = [("Команда проекта", f'{num(inp["C1"]["b"])} FTE × {num(inp["C2"]["b"])} мес.', c["team"]),
                 ("Инфраструктура, данные, ПО", "", c["infra"]), ("Контакт-центр: обучение, скрипты", "", c["cc"]),
                 ("Аудит, комплаенс, ИБ", "", c["audit"]), ("Резерв", pct(inp["C6"]["b"]), c["reserve"])]
    ct = "".join(f'<tr><td>{a}<span class="muted"> {b}</span></td><td class="num">{num(v, 1)}</td></tr>' for a, b, v in cost_rows)
    ct += f'<tr class="total"><td>Итого внедрение (лимит 50)</td><td class="num">{num(c["total"], 1)}</td></tr>'
    ct += f'<tr><td class="muted">Эксплуатация после запуска, в год</td><td class="num muted">{num(c["opex"], 1)}</td></tr>'
    sc = "".join(
        f'<tr{" class=hl" if k == "b" else ""}><td>{n}</td><td class="num">{num(s[k]["net"])}</td>'
        f'<td class="num">{num(s[k]["payback_month"])}-й</td><td class="num">{num(s[k]["cum36"])}</td></tr>'
        for k, n in (("p", "Пессимист"), ("b", "База"), ("o", "Оптимист")))
    t1, t3 = int(inp["T1"]["b"]), int(inp["T3"]["b"])
    stages = [("Проектирование, комплаенс", 1, 1), ("MVP, скоринг в теневом режиме", 1, 3),
              ("Пилот 20% трафика, A/B", 3, 5), ("Масштаб 100%, эксплуатация", 5, 6)]
    cw = 62
    gantt = ['<div style="position:relative">', '<div style="display:flex;margin-left:230px">']
    gantt += [f'<div class="small muted c" style="width:{cw}px;text-align:center">М{m}</div>' for m in range(1, 7)]
    gantt.append("</div>")
    for name, a, b in stages:
        gantt.append(f'<div style="display:flex;align-items:center;height:27px"><div class="small" style="width:230px;padding-right:8px;white-space:nowrap">{name}</div>'
                     f'<div style="position:relative;width:{cw * 6}px;height:16px;background:repeating-linear-gradient(90deg,#f4f6f9 0 {cw - 1}px,#e6e5e0 {cw - 1}px {cw}px)">'
                     f'<div style="position:absolute;left:{(a - 1) * cw + 2}px;width:{(b - a + 1) * cw - 4}px;top:0;bottom:0;background:var(--blue);border-radius:3px"></div></div></div>')
    gantt.append(f'<div class="note" style="margin-top:4px">Эффект: с {t1}-го месяца пилот (20% трафика), с {t3}-го — полный.</div></div>')
    risks = [("Качество и задержка данных", "пауза для догрузки и перепроверки; флаг «данные неполные»"),
             ("Рост трения", "мониторинг доли ограничений и NPS, A/B, откат порогов"),
             ("Дрейф модели, адаптация мошенников", "переобучение по заявлениям ОБДС, теневой режим новых версий"),
             ("Изменения регулирования", "регуляторный слой отдельно; правила 161-ФЗ не обучаются")]
    rt = "".join(f'<tr><td class="b">{a}</td><td>{b}</td></tr>' for a, b in risks)
    title = (f"Внедрение {num(c['total'], 1)} млн ₽ за 6 месяцев; чистый эффект {num(B['net'])} млн ₽ в год, "
             f"окупаемость на {num(B['payback_month'])}-м месяце")
    body = f"""
<div class="row" style="gap:24px">
  <div style="flex:0 0 290px"><div class="h3">Смета внедрения<span class="sub">млн ₽</span></div><table class="tight">{ct}</table></div>
  <div style="flex:0 0 530px"><div class="h3">Эффект для банка, база<span class="sub">млн ₽ в год</span></div>{waterfall(mv)}</div>
  <div style="flex:1"><div class="h3">Сценарии</div>
    <table><tr><th>Сценарий</th><th class="num">В год</th><th class="num">Окупаемость</th><th class="num">36 мес.</th></tr>{sc}</table>
    <div class="note" style="margin-top:6px">Чистый эффект, млн ₽. Возмещения — избежанный расход банка ({pct(inp["A21"]["b"])} сохранённых средств клиентов; сами они не доход банка). Ограничения — экономия на ошибочных ограничениях; расходы — эксплуатация.</div>
  </div>
</div>
<div class="row" style="gap:28px;margin-top:8px">
  <div style="flex:0 0 620px"><div class="h3">Дорожная карта</div>{"".join(gantt)}</div>
  <div style="flex:1"><div class="h3">Риски и меры</div><table class="tight">{rt}</table></div>
</div>
<div class="card key small" style="margin-top:8px;padding:7px 12px"><b class="navy">Рекомендация:</b> запустить В3 с теневого режима и пилота на 20% трафика; решение о масштабировании — по A/B-тесту потерь и доли ограничений.</div>"""
    return slide(7, "Экономика, сроки, риски", title, body,
                 "Расчёт команды на допущениях (гипотезы помечены в Excel): иллюстрация механизма, не прогноз. Чувствительность — Приложение 2.", draft=True)


def s8():
    people = [("Аналитик", "логика решения, KPI и экономика"), ("Разработчик 1", "модель экономики и симуляция"),
              ("Разработчик 2", "процесс, архитектура и визуал")]
    cards = "".join(f"""
<div class="card" style="flex:1;padding:22px">
  <div style="width:84px;height:84px;border-radius:50%;background:var(--soft);border:1.5px dashed var(--muted);display:flex;align-items:center;justify-content:center" class="small muted">фото</div>
  <div style="font-size:22px;font-weight:bold;margin-top:14px">[Имя Фамилия]</div>
  <div class="blue b" style="margin-top:4px">{r}</div>
  <div class="small muted" style="margin-top:2px">{d}</div>
  <div style="margin-top:12px">[Вуз, курс]</div>
  <div class="small" style="margin-top:6px">[Опыт: 1–2 строки — проекты, кейс-чемпионаты, стажировки]</div>
  <div class="small muted" style="margin-top:10px">[контакт]</div>
</div>""" for r, d in people)
    body = f'<div class="row" style="gap:24px;margin-top:20px">{cards}</div>'
    return slide(8, "Резюме команды", "Команда «Омский пряник»", body, "Заполнить данные участников до 1 ноября.")


# ------------------------------------------------------------------ приложения
def appx(n, title, body, foot=""):
    return slide(f"П{n}", f"Приложение {n}", title, body, foot)


def fmt_inp(v, unit):
    if v is None:
        return ""
    if unit == "п.п.":
        return num(v * 100, 1) + " п.п."
    if unit == "%":
        return pct(v, 1).replace(",0%", "%")
    if isinstance(v, float) and not v.is_integer():
        return num(v, 2 if v < 1 else 1)
    return num(v)


def a1(inp):
    keys = ["A1", "A2", "A5", "A7", "A8", "A9a", "A10a", "A10b", "A11a", "A11b", "A12a", "A12b", "A14a", "A14b", "A14c",
            "A15a", "A15b", "A16a", "A17", "A18", "A19", "A20", "A21", "A22", "A23"]
    rows = []
    for k in keys:
        d = inp[k]
        unit = d["unit"] or ""
        if d["p"] is not None:
            val = " / ".join(fmt_inp(d[s], unit) for s in ("p", "b", "o"))
        else:
            val = fmt_inp(d["b"], unit)
        u = "" if unit in ("%", "п.п.") else unit
        st = d["status"] or ""
        stc = ' class="accent"' if st == "гипотеза" else ""
        rows.append(f'<tr><td class="muted">{k}</td><td>{d["name"]}</td><td class="num">{val}{(" " + u) if u else ""}</td>'
                    f'<td>{d["range"] or ""}</td><td{stc}>{st}</td></tr>')
    half = (len(rows) + 1) // 2
    head = '<tr><th>№</th><th>Параметр</th><th class="num">Значение</th><th>Диапазон</th><th>Статус</th></tr>'
    body = (f'<div class="row" style="gap:28px"><div style="flex:1"><table class="tight">{head}{"".join(rows[:half])}</table></div>'
            f'<div style="flex:1"><table class="tight">{head}{"".join(rows[half:])}</table></div></div>'
            '<div class="note" style="margin-top:8px">Три значения через «/» — пессимист / база / оптимист. Полный список с источниками — Excel, лист «Допущения».</div>')
    return appx(1, "Ключевые допущения модели", body, "Статус: «кейс» — материалы кейса, «гипотеза» — допущение команды с диапазоном.")


def a_img(n, title, img, note, foot):
    body = f'<img class="img-full" src="../../out/{img}"><div class="note" style="margin-top:8px">{note}</div>'
    return appx(n, title, body, foot)


def a4(sim):
    rows = "".join(
        f'<tr{" class=hl" if v == "В3" else ""}><td style="white-space:nowrap"><span class="vdot" style="background:{VCOL[v]}"></span>{v}</td>'
        f'<td class="num">{pct(sim[v]["precision"], 1)}</td><td class="num">{pct(sim[v]["recall"], 1)}</td>'
        f'<td class="num">{pct(sim[v]["fpr"], 2)}</td><td class="num">{pct(sim[v]["loss_share"])}</td>'
        f'<td class="num">{pct(sim[v]["restricted_rate_at_equal_loss"], 2)}</td></tr>' for v in VARIANTS)
    body = f"""
<div class="row" style="gap:32px">
  <div style="flex:0 0 500px" class="small">
    <div class="h3">Что смоделировано</div>
    <ul class="clean">
      <li>575 тыс. переводов 30 тыс. синтетических клиентов за 60 дней: возраст, типичная сумма, частота, время суток.</li>
      <li>Мошенничество — 0,3% операций. А «безопасный счёт»: 65% случаев, 70% жертв 55+, звонок перед переводом. Б «захват доступа»: смена устройства или номера, затем перевод.</li>
      <li>Шум: законные операции с похожими признаками — покупка после смены телефона (как у стажёра), крупный перевод 55+ после звонка продавца.</li>
      <li>Задержка до суток: в момент операции видна только часть признаков (смена номера — 50%, входящий звонок — 40%, база ФинЦЕРТ — 70%).</li>
      <li>Регуляторный слой одинаков во всех вариантах; модель его не снимает и на нём не обучается.</li>
      <li>Скоринг — градиентный бустинг; обучение на 36 днях, проверка на следующих 24.</li>
    </ul>
    <div class="h3" style="margin-top:10px">Как читать</div>
    <div>Сигнал — любая мера, кроме «пропустить». Ограничение — пауза, звонок или отказ. Равное трение: у всех вариантов столько же ограничений, сколько у В1.</div>
  </div>
  <div style="flex:1">
    <div class="h3">Результат на тестовой выборке</div>
    <table class="wrap-head"><tr><th></th><th class="num">Точность</th><th class="num">Полнота</th><th class="num">Ложные среди законных</th>
    <th class="num">Потери при равном трении</th><th class="num">Ограничения при потерях 35%</th></tr>{rows}</table>
    <div class="note" style="margin-top:4px">В1 — правила; В2 — ML-скоринг; В3 — адаптивное трение; В4 — внешние участники. Полнота — по сигналам.</div>
    <div class="note" style="margin-top:10px">Точность В3 ниже, чем у В2: сигналов больше, но {pct(sim["В3"]["signals_cleared_share"])} из них — предупреждение или биометрия без остановки платежа.
    На 4 разных случайных выборках порядок вариантов не меняется. Без задержки данных потери В3 были бы ещё ниже: пауза частично компенсирует задержку.</div>
  </div>
</div>"""
    return appx(4, "Симуляция на синтетических данных: как устроена и что показывает", body,
                "Иллюстрация механизма, не прогноз: данных ВТБ в симуляции нет. Код и данные — sim/simulate.py, out/sim_*.csv.")


def a5():
    caps = ["Контекстное предупреждение", "Быстрое подтверждение", "Пауза с объяснением", "«Защита близких»"]
    notes = ["Вопрос про звонок «от банка» прерывает сценарий «безопасный счёт».",
             "Код из SMS можно продиктовать мошеннику, биометрию — нет.",
             "Таймер, причина и две кнопки; «продолжить» — где допускает закон.",
             "Доверенное лицо подтверждает крупные переводы."]
    cells = "".join(f'<div style="flex:1;display:flex;flex-direction:column;align-items:center;gap:8px">'
                    f'{phone_small(i + 1, 0.76)}<div class="b small" style="text-align:center">{caps[i]}</div>'
                    f'<div class="note" style="text-align:center;max-width:250px">{notes[i]}</div></div>' for i in range(4))
    return appx(5, "Макеты экранов: клиент видит меру, соразмерную риску", f'<div class="row" style="gap:18px">{cells}</div>',
                "Нейтральный стиль без фирменных элементов банка; суммы и имена условные.")


def a6():
    rows = [
        ("Скрипт звонка при паузе", "помочь клиенту распознать обман без давления", "кейс высокого риска, тип сценария", "решение клиента: отмена или продолжение, отметка «обман»"),
        ("Методика порогов по сегментам", "единые правила настройки порогов", "KPI, A/B-тесты, сегменты риска", "пороги движка решений по сегментам"),
        ("Регламент разбора и возврата", "сроки и порядок разбора, отделить 161-ФЗ от 115-ФЗ", "обращения, заявления об ОБДС", "решение по кейсу, данные для обучения"),
        ("Тексты предупреждений", "простые формулировки без угроз и юридического языка", "тип сценария, мера", "экранные тексты, SMS, push"),
        ("Обучение контакт-центра", "одинаковая работа операторов с паузами и жалобами", "скрипты, регламенты, разборы ошибок", "обученные операторы, контроль качества"),
        ("Коммуникация «Защиты близких»", "увеличить подключения в сегменте 55+", "клиенты 55+ после паузы или звонка", "предложение сервиса в момент риска"),
    ]
    tr = "".join(f'<tr><td class="b">{a}</td><td>{b}</td><td>{c}</td><td>{d}</td></tr>' for a, b, c, d in rows)
    body = (f'<table><tr><th style="width:260px">Блок</th><th>Назначение</th><th>Вход</th><th>Выход</th></tr>{tr}</table>'
            '<div class="note" style="margin-top:10px">Технические блоки с назначением, входом и выходом — на слайде 5.</div>')
    return appx(6, "Блоки бизнес-процесса: назначение, вход, выход", body, "Внутреннее устройство алгоритмов не раскрываем — по условию кейса.")


def a7():
    src = [
        ("Банк России. Обзор операций без добровольного согласия клиентов за I и II кв. 2026 г.", "cbr.ru/statistics/ib/review_2q_2026"),
        ("Банк России. Обзор операций без добровольного согласия клиентов за 2025 г.", "cbr.ru/analytics/ib/operations_survey/2025"),
        ("Приказ Банка России от 05.11.2025 № ОД-2506 о признаках переводов без добровольного согласия", "cbr.ru/Crosscut/LawActs/File/10123"),
        ("Федеральный закон от 27.06.2011 № 161-ФЗ «О национальной платёжной системе»", ""),
        ("Федеральный закон № 41-ФЗ (с 01.09.2025): «второе лицо» подтверждает операции в течение 12 часов", ""),
        ("Федеральный закон № 210-ФЗ (с 01.03.2027): право банка придержать платёж до 6 часов", ""),
        ("Материалы кейса «Деньги под защитой» (ВТБ, Changellenge, Cup Moscow 2026)", ""),
        ("Модель и симуляция команды: CupMoscow2026-ВТБ-ОмскийПряник.xlsx; sim/simulate.py", ""),
    ]
    li = "".join(f'<li>{a}{"<span class=muted> — " + b + "</span>" if b else ""}</li>' for a, b in src)
    return appx(7, "Источники", f'<ul class="clean" style="font-size:15px;line-height:1.5">{li}</ul>')


# ------------------------------------------------------------------ сборка
def build():
    mv, inp, factors, sim, curves, measures = load()
    slides = [s1(mv), s2(mv, inp), s3(mv, inp, factors, sim), s4(), s5(), s6(mv, sim, curves), s7(mv, inp), s8(),
              a1(inp),
              a_img(2, "Чувствительность: что сильнее всего меняет чистый эффект",
                    "04_chuvstvitelnost_tornado.png",
                    "Каждое допущение меняется на ±30% при прочих равных. Сильнее всего на эффект влияют объём законных переводов и доля ошибочных ограничений после внедрения.",
                    "Расчёт — Excel, лист «Чувствительность»."),
              a_img(3, "Накопленный эффект и окупаемость по сценариям", "05_nakoplennyi_effekt.png",
                    "Затраты внедрения распределены по 6 месяцам; эффект — с пилота (20% трафика), затем полный.",
                    "Расчёт — Excel, листы «Затраты», «Эффект», «Сценарии»."),
              a4(sim),
              a5(), a6(), a7(),
              a_img(8, "Потери против трения: все настройки порогов каждого варианта", "01_poteri_vs_trenie.png",
                    "Каждая линия — варианты порогов одного подхода; точки — рабочие пороги при равном трении.",
                    "Синтетические данные: иллюстрация механизма, не прогноз.")]
    html = f"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8">
<title>CupMoscow2026-ВТБ-ОмскийПряник</title>
<link rel="stylesheet" href="../style.css"></head>
<body>{"".join(slides)}</body></html>"""
    BUILD.mkdir(parents=True, exist_ok=True)
    page = BUILD / "deck.html"
    page.write_text(html, encoding="utf-8")
    print("html:", page)
    if "--html-only" in sys.argv:
        return
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--hide-scrollbars",
                    f"--print-to-pdf={PDF}", "--virtual-time-budget=5000", page.as_uri()],
                   check=True, capture_output=True)
    print(f"pdf: {PDF} ({PDF.stat().st_size / 1e6:.2f} МБ)")


if __name__ == "__main__":
    build()
